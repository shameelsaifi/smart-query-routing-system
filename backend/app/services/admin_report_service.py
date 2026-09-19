from __future__ import annotations

from datetime import (
    datetime,
    timezone,
)
from html import escape
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError

from app.core.database import SessionLocal
from app.services.ticket_access_service import (
    load_ticket_actor,
)


def _load_report_data(
    current_user: dict[str, Any],
) -> tuple[
    dict[str, Any],
    dict[str, Any],
]:
    try:
        with SessionLocal() as db:
            actor = load_ticket_actor(
                db,
                current_user,
                "ADMIN",
            )

            summary = dict(
                db.execute(
                    text(
                        """
                        SELECT
                            COUNT(*) FILTER (
                                WHERE status <> 'DRAFT'
                            )::int
                                AS total_queries,

                            COUNT(*) FILTER (
                                WHERE status NOT IN (
                                    'DRAFT',
                                    'RESOLVED',
                                    'CLOSED'
                                )
                            )::int
                                AS open_queries,

                            COUNT(*) FILTER (
                                WHERE status = 'RESOLVED'
                            )::int
                                AS resolved_queries,

                            COUNT(*) FILTER (
                                WHERE status = 'ESCALATED'
                            )::int
                                AS escalated_queries,

                            COUNT(*) FILTER (
                                WHERE status NOT IN (
                                    'DRAFT',
                                    'RESOLVED',
                                    'CLOSED'
                                )
                                AND sla_due_at IS NOT NULL
                                AND sla_due_at <= NOW()
                            )::int
                                AS overdue_queries,

                            ROUND(
                                COALESCE(
                                    AVG(
                                        EXTRACT(
                                            EPOCH FROM (
                                                resolved_at
                                                -
                                                COALESCE(
                                                    submitted_at,
                                                    created_at
                                                )
                                            )
                                        ) / 3600
                                    ) FILTER (
                                        WHERE status = 'RESOLVED'
                                          AND resolved_at
                                              IS NOT NULL
                                    ),
                                    0
                                )::numeric,
                                2
                            )
                                AS avg_resolution_hours,

                            (
                                SELECT COUNT(*)::int
                                FROM public.users
                                WHERE is_active = TRUE
                            )
                                AS active_users,

                            (
                                SELECT COUNT(*)::int
                                FROM public.escalations
                                WHERE escalation_status =
                                      'OPEN'
                            )
                                AS open_escalations

                        FROM public.tickets
                        """
                    )
                ).mappings().one()
            )

            total_queries = int(
                summary.get(
                    "total_queries",
                    0,
                )
                or 0
            )

            resolved_queries = int(
                summary.get(
                    "resolved_queries",
                    0,
                )
                or 0
            )

            summary[
                "resolution_rate"
            ] = (
                round(
                    (
                        resolved_queries
                        / total_queries
                    )
                    * 100,
                    2,
                )
                if total_queries
                else 0.0
            )

            departments = [
                dict(row)
                for row in db.execute(
                    text(
                        """
                        SELECT
                            d.department_id::text
                                AS department_id,

                            d.department_name,

                            COUNT(t.ticket_id) FILTER (
                                WHERE t.status <> 'DRAFT'
                            )::int
                                AS total_queries,

                            COUNT(t.ticket_id) FILTER (
                                WHERE t.status NOT IN (
                                    'DRAFT',
                                    'RESOLVED',
                                    'CLOSED'
                                )
                            )::int
                                AS active_queries,

                            COUNT(t.ticket_id) FILTER (
                                WHERE t.status =
                                      'ESCALATED'
                            )::int
                                AS escalated_queries,

                            COUNT(t.ticket_id) FILTER (
                                WHERE t.status =
                                      'RESOLVED'
                            )::int
                                AS resolved_queries,

                            COUNT(t.ticket_id) FILTER (
                                WHERE t.status NOT IN (
                                    'DRAFT',
                                    'RESOLVED',
                                    'CLOSED'
                                )
                                AND t.sla_due_at
                                    IS NOT NULL
                                AND t.sla_due_at
                                    <= NOW()
                            )::int
                                AS overdue_queries,

                            ROUND(
                                COALESCE(
                                    AVG(
                                        EXTRACT(
                                            EPOCH FROM (
                                                t.resolved_at
                                                -
                                                COALESCE(
                                                    t.submitted_at,
                                                    t.created_at
                                                )
                                            )
                                        ) / 3600
                                    ) FILTER (
                                        WHERE t.status =
                                              'RESOLVED'
                                          AND t.resolved_at
                                              IS NOT NULL
                                    ),
                                    0
                                )::numeric,
                                2
                            )
                                AS avg_resolution_hours

                        FROM public.departments d

                        LEFT JOIN public.query_assignments qa
                          ON qa.department_id =
                             d.department_id
                         AND qa.is_current = TRUE

                        LEFT JOIN public.tickets t
                          ON t.ticket_id =
                             qa.ticket_id

                        WHERE d.is_active = TRUE

                        GROUP BY
                            d.department_id,
                            d.department_name

                        ORDER BY
                            d.department_name
                        """
                    )
                ).mappings().all()
            ]

            tickets = [
                dict(row)
                for row in db.execute(
                    text(
                        """
                        SELECT
                            t.ticket_number,
                            t.subject,
                            t.source,

                            COALESCE(
                                qc.category_name,
                                t.category
                            )
                                AS category,

                            t.priority,
                            t.status,

                            d.department_name,

                            desk.desk_name,

                            officer.full_name
                                AS assigned_officer,

                            officer.email
                                AS assigned_officer_email,

                            t.submitted_at,
                            t.sla_due_at,
                            t.resolved_at,
                            t.created_at

                        FROM public.tickets t

                        LEFT JOIN
                            public.query_categories qc
                          ON qc.category_id =
                             t.category_id

                        LEFT JOIN LATERAL (
                            SELECT
                                qa.department_id,
                                qa.desk_id,
                                qa.assigned_user_id

                            FROM public.query_assignments qa

                            WHERE qa.ticket_id =
                                  t.ticket_id
                              AND qa.is_current = TRUE

                            ORDER BY
                                qa.assigned_at DESC

                            LIMIT 1
                        ) current_assignment
                          ON TRUE

                        LEFT JOIN public.departments d
                          ON d.department_id =
                             current_assignment
                                 .department_id

                        LEFT JOIN public.accounts_desks desk
                          ON desk.desk_id =
                             COALESCE(
                                 t.routed_desk_id,
                                 current_assignment.desk_id
                             )

                        LEFT JOIN public.users officer
                          ON officer.user_id =
                             COALESCE(
                                 t.assigned_officer_id,
                                 current_assignment
                                     .assigned_user_id
                             )

                        WHERE t.status <> 'DRAFT'

                        ORDER BY
                            COALESCE(
                                t.submitted_at,
                                t.created_at
                            ) DESC,
                            t.ticket_number DESC
                        """
                    )
                ).mappings().all()
            ]

            report_data = {
                "generated_at":
                    datetime.now(
                        timezone.utc
                    ),

                "summary":
                    summary,

                "departments":
                    departments,

                "tickets":
                    tickets,
            }

            return actor, report_data

    except SQLAlchemyError as exc:
        raise HTTPException(
            status_code=503,
            detail=(
                "Report data is temporarily "
                "unavailable."
            ),
        ) from exc


def _record_export(
    actor_id: str,
    report_format: str,
    ticket_count: int,
) -> None:
    try:
        with SessionLocal.begin() as db:
            db.execute(
                text(
                    """
                    INSERT INTO public.audit_logs (
                        actor_user_id,
                        actor_service,
                        action,
                        entity_type,
                        entity_id,
                        outcome,
                        old_values,
                        new_values,
                        details
                    )
                    VALUES (
                        CAST(
                            :actor_id
                            AS UUID
                        ),
                        NULL,
                        'ADMIN_REPORT_EXPORTED',
                        'REPORT',
                        NULL,
                        'SUCCESS',
                        NULL,
                        NULL,
                        jsonb_build_object(
                            'format',
                            CAST(
                                :report_format
                                AS TEXT
                            ),
                            'ticket_count',
                            CAST(
                                :ticket_count
                                AS INTEGER
                            )
                        )
                    )
                    """
                ),
                {
                    "actor_id":
                        actor_id,

                    "report_format":
                        report_format,

                    "ticket_count":
                        ticket_count,
                },
            )

    except SQLAlchemyError as exc:
        raise HTTPException(
            status_code=503,
            detail=(
                "Report audit could not "
                "be recorded."
            ),
        ) from exc


def _excel_datetime(
    value: Any,
) -> Any:
    if not isinstance(
        value,
        datetime,
    ):
        return value

    if value.tzinfo is not None:
        value = (
            value.astimezone(
                timezone.utc
            )
            .replace(
                tzinfo=None
            )
        )

    return value


def _apply_excel_header(
    worksheet: Any,
    row_number: int,
    column_count: int,
) -> None:
    fill = PatternFill(
        fill_type="solid",
        fgColor="1E3A8A",
    )

    font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin = Side(
        style="thin",
        color="CBD5E1",
    )

    for column in range(
        1,
        column_count + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column,
        )

        cell.fill = fill
        cell.font = font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin,
        )


def _fit_excel_columns(
    worksheet: Any,
    maximum: int = 42,
) -> None:
    for column_cells in worksheet.columns:
        length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            length = max(
                length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                length + 2,
                11,
            ),
            maximum,
        )


def build_admin_excel_report(
    current_user: dict[str, Any],
) -> dict[str, Any]:
    actor, data = _load_report_data(
        current_user
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_sheet.merge_cells(
        "A1:D1"
    )

    summary_sheet["A1"] = (
        "SmartQuery Administrative Report"
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )

    summary_sheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="0F172A",
    )

    summary_sheet["A1"].alignment = (
        Alignment(
            horizontal="center",
            vertical="center",
        )
    )

    summary_sheet.row_dimensions[
        1
    ].height = 28

    summary_sheet["A3"] = (
        "Generated At (UTC)"
    )

    summary_sheet["B3"] = (
        _excel_datetime(
            data["generated_at"]
        )
    )

    summary_sheet["B3"].number_format = (
        "yyyy-mm-dd hh:mm:ss"
    )

    summary = data["summary"]

    summary_rows = [
        (
            "Total Queries",
            summary["total_queries"],
        ),
        (
            "Open Queries",
            summary["open_queries"],
        ),
        (
            "Resolved Queries",
            summary["resolved_queries"],
        ),
        (
            "Escalated Queries",
            summary["escalated_queries"],
        ),
        (
            "Open Escalations",
            summary["open_escalations"],
        ),
        (
            "Overdue Queries",
            summary["overdue_queries"],
        ),
        (
            "Active Users",
            summary["active_users"],
        ),
        (
            "Resolution Rate",
            f"{summary['resolution_rate']}%",
        ),
        (
            "Average Resolution Hours",
            summary[
                "avg_resolution_hours"
            ],
        ),
    ]

    summary_sheet.append([])
    summary_sheet.append(
        [
            "Metric",
            "Value",
        ]
    )

    _apply_excel_header(
        summary_sheet,
        5,
        2,
    )

    for metric, value in summary_rows:
        summary_sheet.append(
            [
                metric,
                value,
            ]
        )

    _fit_excel_columns(
        summary_sheet
    )

    summary_sheet.freeze_panes = "A6"


    department_sheet = (
        workbook.create_sheet(
            "Departments"
        )
    )

    department_headers = [
        "Department",
        "Total Queries",
        "Active Queries",
        "Escalated",
        "Overdue",
        "Resolved",
        "Avg Resolution Hours",
    ]

    department_sheet.append(
        department_headers
    )

    _apply_excel_header(
        department_sheet,
        1,
        len(department_headers),
    )

    for department in data[
        "departments"
    ]:
        department_sheet.append(
            [
                department[
                    "department_name"
                ],
                department[
                    "total_queries"
                ],
                department[
                    "active_queries"
                ],
                department[
                    "escalated_queries"
                ],
                department[
                    "overdue_queries"
                ],
                department[
                    "resolved_queries"
                ],
                department[
                    "avg_resolution_hours"
                ],
            ]
        )

    department_sheet.freeze_panes = "A2"
    department_sheet.auto_filter.ref = (
        department_sheet.dimensions
    )

    _fit_excel_columns(
        department_sheet
    )


    ticket_sheet = workbook.create_sheet(
        "Tickets"
    )

    ticket_headers = [
        "Ticket #",
        "Subject",
        "Source",
        "Category",
        "Priority",
        "Status",
        "Department",
        "Desk",
        "Assigned Officer",
        "Officer Email",
        "Submitted At (UTC)",
        "SLA Due (UTC)",
        "Resolved At (UTC)",
    ]

    ticket_sheet.append(
        ticket_headers
    )

    _apply_excel_header(
        ticket_sheet,
        1,
        len(ticket_headers),
    )

    for ticket in data["tickets"]:
        ticket_sheet.append(
            [
                ticket["ticket_number"],
                ticket["subject"],
                ticket["source"],
                ticket["category"],
                ticket["priority"],
                ticket["status"],
                ticket["department_name"],
                ticket["desk_name"],
                ticket["assigned_officer"],
                ticket[
                    "assigned_officer_email"
                ],
                _excel_datetime(
                    ticket[
                        "submitted_at"
                    ]
                ),
                _excel_datetime(
                    ticket[
                        "sla_due_at"
                    ]
                ),
                _excel_datetime(
                    ticket[
                        "resolved_at"
                    ]
                ),
            ]
        )

    for row in ticket_sheet.iter_rows(
        min_row=2,
    ):
        for column_index in (
            11,
            12,
            13,
        ):
            row[
                column_index - 1
            ].number_format = (
                "yyyy-mm-dd hh:mm:ss"
            )

    ticket_sheet.freeze_panes = "A2"
    ticket_sheet.auto_filter.ref = (
        ticket_sheet.dimensions
    )

    _fit_excel_columns(
        ticket_sheet,
        maximum=40,
    )

    ticket_sheet.column_dimensions[
        "B"
    ].width = 40

    ticket_sheet.column_dimensions[
        "J"
    ].width = 32

    for row in ticket_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    output = BytesIO()

    workbook.save(
        output
    )

    content = output.getvalue()

    _record_export(
        actor["user_id"],
        "XLSX",
        len(data["tickets"]),
    )

    timestamp = (
        data["generated_at"]
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )

    return {
        "content": content,

        "filename": (
            "smartquery_admin_report_"
            f"{timestamp}.xlsx"
        ),
    }


def _pdf_text(
    value: Any,
) -> str:
    if value is None:
        return "-"

    return escape(
        str(value)
    )


def _pdf_datetime(
    value: Any,
) -> str:
    if not isinstance(
        value,
        datetime,
    ):
        return "-"

    if value.tzinfo is not None:
        value = value.astimezone(
            timezone.utc
        )

    return value.strftime(
        "%Y-%m-%d %H:%M"
    )


def _pdf_footer(
    canvas: Any,
    document: Any,
) -> None:
    canvas.saveState()

    canvas.setFont(
        "Helvetica",
        7,
    )

    canvas.setFillColor(
        colors.grey
    )

    canvas.drawString(
        12 * mm,
        7 * mm,
        "SmartQuery Administrative Report",
    )

    canvas.drawRightString(
        landscape(A4)[0]
        - 12 * mm,
        7 * mm,
        f"Page {document.page}",
    )

    canvas.restoreState()


def build_admin_pdf_report(
    current_user: dict[str, Any],
) -> dict[str, Any]:
    actor, data = _load_report_data(
        current_user
    )

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=(
            "SmartQuery "
            "Administrative Report"
        ),
        author="SmartQuery",
    )

    sample_styles = (
        getSampleStyleSheet()
    )

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=sample_styles[
            "Title"
        ],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        textColor=colors.HexColor(
            "#0F172A"
        ),
        spaceAfter=8,
    )

    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=sample_styles[
            "Heading2"
        ],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor(
            "#1E3A8A"
        ),
        spaceBefore=8,
        spaceAfter=6,
    )

    cell_style = ParagraphStyle(
        "ReportCell",
        parent=sample_styles[
            "BodyText"
        ],
        fontName="Helvetica",
        fontSize=6.3,
        leading=7.5,
    )

    story: list[Any] = []

    story.append(
        Paragraph(
            "SmartQuery Administrative Report",
            title_style,
        )
    )

    story.append(
        Paragraph(
            (
                "Generated at "
                f"{_pdf_datetime(data['generated_at'])} UTC"
            ),
            cell_style,
        )
    )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )


    summary = data["summary"]

    summary_table_data = [
        [
            "Metric",
            "Value",
            "Metric",
            "Value",
        ],
        [
            "Total Queries",
            summary["total_queries"],
            "Open Queries",
            summary["open_queries"],
        ],
        [
            "Resolved",
            summary["resolved_queries"],
            "Escalated",
            summary["escalated_queries"],
        ],
        [
            "Open Escalations",
            summary["open_escalations"],
            "Overdue",
            summary["overdue_queries"],
        ],
        [
            "Resolution Rate",
            f"{summary['resolution_rate']}%",
            "Avg. Resolution Hours",
            summary[
                "avg_resolution_hours"
            ],
        ],
        [
            "Active Users",
            summary["active_users"],
            "",
            "",
        ],
    ]

    summary_table = Table(
        summary_table_data,
        colWidths=[
            48 * mm,
            24 * mm,
            48 * mm,
            24 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#1E3A8A"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (0, 1),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 1),
                    (2, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor(
                        "#CBD5E1"
                    ),
                ),
                (
                    "BACKGROUND",
                    (0, 1),
                    (-1, -1),
                    colors.HexColor(
                        "#F8FAFC"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    story.append(
        Paragraph(
            "Department Performance",
            heading_style,
        )
    )


    department_data = [
        [
            "Department",
            "Total",
            "Active",
            "Escalated",
            "Overdue",
            "Resolved",
            "Avg Hrs",
        ]
    ]

    for department in data[
        "departments"
    ]:
        department_data.append(
            [
                _pdf_text(
                    department[
                        "department_name"
                    ]
                ),
                department[
                    "total_queries"
                ],
                department[
                    "active_queries"
                ],
                department[
                    "escalated_queries"
                ],
                department[
                    "overdue_queries"
                ],
                department[
                    "resolved_queries"
                ],
                department[
                    "avg_resolution_hours"
                ],
            ]
        )

    department_table = Table(
        department_data,
        repeatRows=1,
        colWidths=[
            65 * mm,
            22 * mm,
            22 * mm,
            22 * mm,
            22 * mm,
            22 * mm,
            25 * mm,
        ],
    )

    department_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#1E3A8A"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor(
                        "#CBD5E1"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ALIGN",
                    (1, 1),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#F8FAFC"
                        ),
                    ],
                ),
            ]
        )
    )

    story.append(
        department_table
    )

    story.append(
        PageBreak()
    )

    story.append(
        Paragraph(
            "Detailed Query Report",
            heading_style,
        )
    )


    ticket_data = [
        [
            "Ticket",
            "Subject",
            "Source",
            "Category",
            "Priority",
            "Status",
            "Department",
            "Desk",
            "Officer",
            "Submitted",
            "SLA Due",
            "Resolved",
        ]
    ]

    for ticket in data["tickets"]:
        subject = (
            str(
                ticket["subject"]
                or ""
            )
        )

        if len(subject) > 75:
            subject = (
                subject[:72]
                + "..."
            )

        ticket_data.append(
            [
                Paragraph(
                    _pdf_text(
                        ticket[
                            "ticket_number"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        subject
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket["source"]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket[
                            "category"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket[
                            "priority"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket["status"]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket[
                            "department_name"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket[
                            "desk_name"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_text(
                        ticket[
                            "assigned_officer"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_datetime(
                        ticket[
                            "submitted_at"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_datetime(
                        ticket[
                            "sla_due_at"
                        ]
                    ),
                    cell_style,
                ),

                Paragraph(
                    _pdf_datetime(
                        ticket[
                            "resolved_at"
                        ]
                    ),
                    cell_style,
                ),
            ]
        )

    ticket_table = Table(
        ticket_data,
        repeatRows=1,
        colWidths=[
            18 * mm,
            40 * mm,
            13 * mm,
            25 * mm,
            15 * mm,
            23 * mm,
            27 * mm,
            24 * mm,
            28 * mm,
            24 * mm,
            24 * mm,
            24 * mm,
        ],
    )

    ticket_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#0F172A"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, 0),
                    6.3,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.25,
                    colors.HexColor(
                        "#CBD5E1"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#F8FAFC"
                        ),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    3,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    3,
                ),
            ]
        )
    )

    story.append(
        ticket_table
    )

    document.build(
        story,
        onFirstPage=_pdf_footer,
        onLaterPages=_pdf_footer,
    )

    content = output.getvalue()

    _record_export(
        actor["user_id"],
        "PDF",
        len(data["tickets"]),
    )

    timestamp = (
        data["generated_at"]
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )

    return {
        "content": content,

        "filename": (
            "smartquery_admin_report_"
            f"{timestamp}.pdf"
        ),
    }