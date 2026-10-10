"""Role-scoped Staff/Instructor query reports.

Uses the existing get_assigned_tickets access path, not the Admin-wide report
query. Never use build_admin_* here: those reports intentionally contain data
from every department.
"""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone, timedelta
from html import escape
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError

from app.core.database import SessionLocal
from app.services.ticket_access_service import load_ticket_actor
from app.services.ticket_service import get_assigned_tickets


def _as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        date = value
    elif isinstance(value, str) and value:
        try:
            date = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
    else:
        return None
    return date.replace(tzinfo=timezone.utc) if date.tzinfo is None else date.astimezone(timezone.utc)


def _excel_safe(value: Any) -> Any:
    """Prevent spreadsheet formula injection in student-controlled text."""
    if isinstance(value, datetime):
        return _as_datetime(value).replace(tzinfo=None)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value if value is not None else ""


def _ticket_num(ticket: dict[str, Any]) -> str:
    return str(ticket.get("ticket_number") or ticket.get("ticket_code") or ticket.get("ticket_id") or "")


def _filter_tickets(
    items: list[dict[str, Any]],
    *,
    search: str,
    source: str,
    priority: str,
    status: str,
    sla_only: bool,
) -> list[dict[str, Any]]:
    search = search.strip().casefold()
    now = datetime.now(timezone.utc)
    selected = []
    for item in items:
        state = str(item.get("status") or "").upper()
        item_source = str(item.get("source") or "").upper()
        item_priority = str(item.get("priority") or "").upper()
        haystack = " ".join(str(part or "") for part in (
            _ticket_num(item), item.get("student_name"),
            item.get("subject"), item.get("category"),
        )).casefold()
        if search and search not in haystack:
            continue
        if source != "ALL" and item_source != source:
            continue
        if priority != "ALL" and item_priority != priority:
            continue
        if status == "OPEN" and state in {"RESOLVED", "CLOSED"}:
            continue
        if status not in {"OPEN", "ALL"} and state != status:
            continue
        if sla_only:
            due = _as_datetime(item.get("sla_due_at"))
            if due is None or due - now > timedelta(hours=4):
                continue
        selected.append(item)
    return selected


def _load_data(current_user: dict[str, Any], filters: dict[str, Any]) -> dict[str, Any]:
    # Defense in depth: check the live application account, not just the JWT role.
    try:
        with SessionLocal() as db:
            actor = load_ticket_actor(
                db, current_user, "DEPARTMENT_STAFF", "INSTRUCTOR",
            )
    except SQLAlchemyError as exc:
        raise HTTPException(status_code=503, detail="Report access is temporarily unavailable.") from exc

    # Crucial: this is the SAME authorized reader as GET /tickets/assigned-to-me.
    # It must not be replaced by the Admin-wide unscoped report SQL.
    raw = get_assigned_tickets(current_user)
    items = [dict(item) for item in raw]
    selected = _filter_tickets(items, **filters)
    resolved = [item for item in selected if str(item.get("status") or "").upper() == "RESOLVED"]
    durations = []
    for item in resolved:
        submitted = _as_datetime(item.get("submitted_at") or item.get("created_at"))
        ended = _as_datetime(item.get("resolved_at"))
        if submitted and ended and ended >= submitted:
            durations.append((ended - submitted).total_seconds() / 3600)
    states = Counter(str(item.get("status") or "UNKNOWN").upper() for item in selected)
    categories = Counter(str(item.get("category") or "Uncategorized") for item in selected)
    now = datetime.now(timezone.utc)
    overdue = sum(
        1 for item in selected
        if str(item.get("status") or "").upper() not in {"RESOLVED", "CLOSED"}
        and (due := _as_datetime(item.get("sla_due_at"))) is not None
        and due <= now
    )
    return {
        "actor": actor,
        "generated_at": now,
        "filters": filters,
        "tickets": selected,
        "summary": {
            "assigned_total": len(selected),
            "active": sum(count for state, count in states.items() if state not in {"RESOLVED", "CLOSED"}),
            "resolved": states["RESOLVED"],
            "closed": states["CLOSED"],
            "escalated": states["ESCALATED"],
            "overdue": overdue,
            "avg_resolution_hours": round(sum(durations) / len(durations), 2) if durations else None,
        },
        "categories": categories.most_common(),
    }


def _record_export(actor_id: str, report_format: str, count: int) -> None:
    # Match the proven audit schema used by admin_report_service.py.
    try:
        with SessionLocal.begin() as db:
            db.execute(
                text("""
                    INSERT INTO public.audit_logs (
                        actor_user_id, actor_service, action,
                        entity_type, entity_id, outcome,
                        old_values, new_values, details
                    ) VALUES (
                        CAST(:actor_id AS UUID), NULL,
                        'STAFF_REPORT_EXPORTED', 'REPORT', NULL,
                        'SUCCESS', NULL, NULL,
                        jsonb_build_object('format', CAST(:format AS TEXT),
                                           'ticket_count', CAST(:count AS INTEGER),
                                           'scope', 'ASSIGNED_TO_ME')
                    )
                """),
                {"actor_id": actor_id, "format": report_format, "count": count},
            )
    except SQLAlchemyError as exc:
        raise HTTPException(status_code=503, detail="The report audit could not be saved.") from exc


def _timestamp(data: dict[str, Any]) -> str:
    return data["generated_at"].strftime("%Y%m%d_%H%M%S")


def build_staff_excel_report(
    current_user: dict[str, Any], filters: dict[str, Any],
) -> dict[str, Any]:
    data = _load_data(current_user, filters)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["SmartQuery - Authorized Staff Queue Report"])
    summary.append(["Generated (UTC)", data["generated_at"].strftime("%Y-%m-%d %H:%M:%S")])
    summary.append(["Staff", data["actor"].get("full_name", "")])
    summary.append(["Scope", "Only assigned, authorized queries"])
    summary.append([])
    summary.append(["Metric", "Value"])
    titles = {
        "assigned_total": "Exported queries", "active": "Active",
        "resolved": "Resolved", "closed": "Closed",
        "escalated": "Escalated", "overdue": "Overdue",
        "avg_resolution_hours": "Average resolution hours (resolved with timestamps)",
    }
    for key, label in titles.items():
        summary.append([label, data["summary"].get(key) if data["summary"].get(key) is not None else "N/A"])

    categories = workbook.create_sheet("Categories")
    categories.append(["Category", "Assigned queries"])
    for category, count in data["categories"]:
        categories.append([_excel_safe(category), count])

    tickets = workbook.create_sheet("Tickets")
    headers = [
        "Ticket #", "Student", "Subject", "Source", "Category",
        "Priority", "Status", "Assigned To", "Submitted (UTC)",
        "SLA Due (UTC)", "Resolved (UTC)",
    ]
    tickets.append(headers)
    for row in data["tickets"]:
        tickets.append([
            _excel_safe(_ticket_num(row)),
            _excel_safe(row.get("student_name")),
            _excel_safe(row.get("subject")),
            _excel_safe(row.get("source")),
            _excel_safe(row.get("category")),
            _excel_safe(row.get("priority")),
            _excel_safe(row.get("status")),
            _excel_safe(row.get("assigned_user_name") or row.get("assignee_name")),
            _excel_safe(_as_datetime(row.get("submitted_at") or row.get("created_at"))),
            _excel_safe(_as_datetime(row.get("sla_due_at"))),
            _excel_safe(_as_datetime(row.get("resolved_at"))),
        ])
    for cell_row in tickets.iter_rows(min_row=2):
        for index in (8, 9, 10):
            cell_row[index].number_format = "yyyy-mm-dd hh:mm"

    for sheet, header_row in ((summary, 6), (categories, 1), (tickets, 1)):
        sheet.freeze_panes = f"A{header_row + 1}"
        for cell in sheet[header_row]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
            cell.alignment = Alignment(vertical="center")
        for cells in sheet.columns:
            letter = get_column_letter(cells[0].column)
            max_len = max((len(str(c.value)) for c in cells if c.value is not None), default=12)
            sheet.column_dimensions[letter].width = min(48, max(13, max_len + 2))
    tickets.auto_filter.ref = tickets.dimensions
    summary.column_dimensions["A"].width = 53

    output = BytesIO()
    workbook.save(output)
    # Audit before returning the file to the authenticated user.
    _record_export(str(data["actor"]["user_id"]), "XLSX", len(data["tickets"]))
    return {"content": output.getvalue(), "filename": f"smartquery_staff_queue_{_timestamp(data)}.xlsx"}


def _para(value: Any, style: ParagraphStyle) -> Paragraph:
    text_value = "—" if value is None or value == "" else str(value)
    return Paragraph(escape(text_value).replace("\n", "<br/>"), style)


def build_staff_pdf_report(
    current_user: dict[str, Any], filters: dict[str, Any],
) -> dict[str, Any]:
    data = _load_data(current_user, filters)
    output = BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=11 * mm,
        rightMargin=11 * mm, topMargin=14 * mm,
        bottomMargin=14 * mm, title="SmartQuery Staff Queue Report",
    )
    samples = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "StaffTitle", parent=samples["Title"], fontSize=17,
        leading=21, textColor=colors.HexColor("#0F172A"),
    )
    small = ParagraphStyle("StaffSmall", parent=samples["BodyText"], fontSize=7, leading=9)
    table_header = ParagraphStyle("StaffHeader", parent=small, textColor=colors.white, fontName="Helvetica-Bold")
    heading = ParagraphStyle("StaffHeading", parent=samples["Heading2"], fontSize=11,
                             textColor=colors.HexColor("#1E3A8A"))
    story = [
        Paragraph("SmartQuery | Staff Queue Report", title_style),
        _para("Generated (UTC): " + data["generated_at"].strftime("%Y-%m-%d %H:%M"), small),
        _para("Scope: currently assigned queries visible to the authenticated staff member", small),
        Spacer(1, 5 * mm), Paragraph("Queue summary", heading),
    ]
    summary = data["summary"]
    summary_rows = [
        ["Exported", "Active", "Resolved", "Escalated", "Overdue", "Avg resolution hours"],
        [summary["assigned_total"], summary["active"], summary["resolved"],
         summary["escalated"], summary["overdue"],
         summary["avg_resolution_hours"] if summary["avg_resolution_hours"] is not None else "N/A"],
    ]
    report_rows = [
        ["Ticket", "Student", "Subject", "Source", "Category", "Priority", "Status", "Submitted (UTC)"],
    ]
    for row in data["tickets"]:
        timestamp = _as_datetime(row.get("submitted_at") or row.get("created_at"))
        report_rows.append([
            _ticket_num(row), row.get("student_name"), row.get("subject"),
            row.get("source"), row.get("category"), row.get("priority"),
            row.get("status"), timestamp.strftime("%Y-%m-%d %H:%M") if timestamp else "",
        ])
    if len(report_rows) == 1:
        report_rows.append(["No matching assigned queries", "", "", "", "", "", "", ""])

    def make_table(rows: list[list[Any]], widths: list[Any]) -> Table:
        table = Table([[_para(value, table_header if index == 0 else small) for value in row] for index, row in enumerate(rows)],
                      colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return table

    story.append(make_table(summary_rows, [34 * mm, 30 * mm, 30 * mm, 34 * mm, 30 * mm, 48 * mm]))
    story.extend([Spacer(1, 5 * mm), Paragraph("Authorized query details", heading),
                  make_table(report_rows, [23 * mm, 32 * mm, 64 * mm, 18 * mm,
                                           38 * mm, 20 * mm, 27 * mm, 51 * mm])])

    def footer(canvas: Any, report_doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(11 * mm, 8 * mm, "SmartQuery — Authorized staff export")
        canvas.drawRightString(landscape(A4)[0] - 11 * mm, 8 * mm,
                               f"Page {report_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    _record_export(str(data["actor"]["user_id"]), "PDF", len(data["tickets"]))
    return {"content": output.getvalue(), "filename": f"smartquery_staff_queue_{_timestamp(data)}.pdf"}
